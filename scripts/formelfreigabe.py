#!/usr/bin/env python3
"""
Formelfreigabe über Excel.

  python3 scripts/formelfreigabe.py export [--out DATEI.xlsx] [--offen] [--modul MODUL[,MODUL] | --modul werkbank|fachtools]
      Erzeugt die Freigabe-Excel aus dem aktuellen Rechenkern (alle Formeln aus scripts/katalog.ts,
      gleiche Gruppierung wie FORMELKATALOG.md, mit nachrechenbarem Beispiel). Bereits vorhandene
      Freigaben aus formelstatus.json werden vorbefüllt. --modul (ab 0.10.0) beschränkt auf die Formeln der
      genannten Module bzw. der Module eines Werkzeugs (src/modules.json); der Import übernimmt dann nur diese
      Zeilen, alle anderen Formeln behalten ihren Status.

  python3 scripts/formelfreigabe.py import DATEI.xlsx [--dry-run] [--ergaenze-datum TT.MM.JJJJ --ergaenze-pruefer NAME]
      Liest die ausgefüllte Excel, prüft die Einträge und schreibt formelstatus.json.
      Danach wird FORMELKATALOG.md neu erzeugt. Abgelehnte Formeln mit Kommentar werden aufgelistet.
      --ergaenze-*: nur für Zeilen »freigegeben« ohne Datum bzw. Prüfer, wenn der Prüfer die Freigabe
      außerhalb der Excel bestätigt hat. Jede ergänzte Zeile bekommt einen Vermerk im Kommentar.

Regeln beim Import:
  - Die Excel ist vollständig maßgeblich: was dort "offen" ist, ist danach offen.
  - "freigegeben" braucht Datum und Prüfer, sonst wird der Eintrag nicht übernommen.
  - Ein Kommentar aus einer früheren Runde, der bei einer Freigabe unverändert in der Zeile steht, wird als
    "erledigt" gespeichert (Verlauf), nicht als offener Kommentar.
  - "abgelehnt" sollte einen Kommentar haben (Warnung, wird trotzdem übernommen).
  - Hat sich eine Formel seit dem Export geändert, wird ihre Zeile nicht übernommen (neu exportieren).
  - Eine Freigabe speichert den Formeltext. Ändert sich die Formel später im Kern, zeigt der Katalog
    "geändert, neu prüfen".

Abhängigkeiten: openpyxl, Node mit tsx (npm ci).
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
STATUS_FILE = ROOT / "formelstatus.json"
SHEET = "Freigabe"
STATES = ("offen", "freigegeben", "abgelehnt")

# Spalten: (Kopf, Breite, Schlüssel). Eingabespalten sind gelb.
COLS = [
    ("ID", 11, "id"),
    ("Bereich", 16, "groupTitle"),
    ("Größe", 30, "title"),
    ("Formel", 40, "formula"),
    ("Beispiel eingesetzt", 40, "exSubst"),
    ("Ergebnis", 13, "exValue"),
    ("Einheit", 8, "unit"),
    ("Quelle", 30, "ref"),
    ("Beispielfall", 44, "exCase"),
    ("Freigabe", 13, "in_status"),
    ("Datum", 12, "in_datum"),
    ("Prüfer", 16, "in_durch"),
    ("Kommentar / Verbesserung", 50, "in_kommentar"),
    ("Formel beim Export (nicht ändern)", 30, "formula_stamp"),
]
C = {key: i + 1 for i, (_, _, key) in enumerate(COLS)}
LAST = len(COLS)
INPUT_KEYS = ("in_status", "in_datum", "in_durch", "in_kommentar")

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_MUTED = Font(name=FONT, size=9, color="595959")
FILL_HEAD = PatternFill("solid", start_color="1F3A5F")
FILL_GROUP = PatternFill("solid", start_color="D9E1EA")
FILL_INPUT = PatternFill("solid", start_color="FFF2CC")
FILL_OK = PatternFill("solid", start_color="C6EFCE")
FILL_BAD = PatternFill("solid", start_color="FFC7CE")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def load_catalog() -> dict:
    """Katalog aus dem TypeScript-Kern holen (eine Quelle für alle Formeln)."""
    res = subprocess.run(["npx", "--no-install", "tsx", "scripts/katalog.ts", "--json"], cwd=ROOT, capture_output=True, text=True)
    if res.returncode != 0:
        sys.exit(f"Katalog konnte nicht erzeugt werden:\n{res.stderr}")
    return json.loads(res.stdout)


def sig(v: float, n: int = 6) -> float:
    return float(f"{v:.{n}g}") if isinstance(v, (int, float)) else v


def parse_date(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(s, f).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"Datum nicht lesbar: {s!r} (erwartet TT.MM.JJJJ)")


# --------------------------------------------------------------------------------------------- Export

def resolve_modules(spec: str | None) -> list[str] | None:
    """--modul: Modulnamen (Komma) oder ein Werkzeug aus src/modules.json."""
    if not spec:
        return None
    reg = json.loads((ROOT / "src" / "modules.json").read_text(encoding="utf-8"))
    if spec in reg["tools"]:
        return list(reg["tools"][spec]["modules"])
    mods = [m.strip() for m in spec.split(",") if m.strip()]
    bad = [m for m in mods if m not in reg["modules"]]
    if bad:
        sys.exit(f"Modul unbekannt: {', '.join(bad)} (bekannt: {', '.join(reg['modules'])}; Werkzeuge: {', '.join(reg['tools'])})")
    return mods


def export(out: Path, only_open: bool = False, modules: list[str] | None = None) -> None:
    cat = load_catalog()
    entries, status = cat["entries"], cat["status"]
    if modules is not None:
        entries = [e for e in entries if e["module"] in modules]
        cat = {**cat, "modules": modules}
    if only_open:
        # nur Formeln ohne gültige Freigabe: offen, abgelehnt oder seit der Freigabe geändert
        def approved(e: dict) -> bool:
            st = status.get(e["id"])
            return bool(st) and st.get("status") == "freigegeben" and st.get("formel") in (None, e["formula"])
        entries = [e for e in entries if not approved(e)]
    if not entries:
        scope = f"Module {', '.join(modules)}" if modules is not None else "Kern"
        sys.exit(f"{scope}: keine Formeln{' ohne gültige Freigabe' if only_open else ''}, keine Excel erzeugt")
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    write_sheet(ws, cat, entries, status)
    write_guide(wb.create_sheet("Anleitung"), cat)
    wb.active = 0
    wb.properties.title = f"Formelfreigabe eng-core {cat['coreVersion']}" + (f" ({', '.join(cat['modules'])})" if cat.get("modules") else "")
    wb.properties.subject = f"{len(entries)} Formeln, exportiert {dt.date.today().isoformat()}"
    wb.properties.creator = "eng-core scripts/formelfreigabe.py"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"{out.relative_to(ROOT) if out.is_relative_to(ROOT) else out}: {len(entries)} Formeln exportiert")


def write_sheet(ws, cat: dict, entries: list, status: dict) -> None:
    today = dt.date.today().strftime("%d.%m.%Y")
    ws["A1"] = f"Formelfreigabe eng-core {cat['coreVersion']}" + (f" · Module {', '.join(cat['modules'])}" if cat.get("modules") else "")
    ws["A1"].font = F_TITLE
    reg = json.loads((ROOT / "src" / "modules.json").read_text(encoding="utf-8"))["modules"]
    uses_fd = any(reg.get(e["module"], {}).get("fd") for e in entries)
    fd = f" · Fachdaten {cat.get('fdVersion') or '–'}" if uses_fd else ""
    ws["A2"] = f"Datenstand Werkstoffdaten {cat['dataVersion']}{fd} · exportiert am {today} · nur gelbe Spalten ausfüllen, Anleitung im zweiten Blatt"
    ws["A2"].font = F_MUTED
    head_row, first = 5, 6

    # Zähler (Formeln, damit sie beim Ausfüllen mitlaufen)
    labels = [("Formeln", None), ("freigegeben", "freigegeben"), ("abgelehnt", "abgelehnt"), ("offen", "offen")]
    status_col = get_column_letter(C["in_status"])
    n_last = first + len(entries) + 20
    for i, (lab, key) in enumerate(labels):
        col = 2 * i + 1
        ws.cell(row=3, column=col, value=lab).font = F_BOLD
        rng_s = f"${status_col}${first}:${status_col}${n_last}"
        stamp_col = get_column_letter(C["formula_stamp"])
        f = f"=COUNTA({stamp_col}{first}:{stamp_col}{n_last})" if key is None else f'=COUNTIF({rng_s},"{key}")'
        c = ws.cell(row=3, column=col + 1, value=f)
        c.font = F_BOLD
        c.alignment = Alignment(horizontal="left")

    for i, (name, width, _) in enumerate(COLS, start=1):
        c = ws.cell(row=head_row, column=i, value=name)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, Alignment(wrap_text=True, vertical="center"), BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[head_row].height = 28

    r, last_group = first, None
    for e in entries:
        if e["group"] != last_group:
            c = ws.cell(row=r, column=1, value=e["groupTitle"])
            c.font = F_BOLD
            for j in range(1, LAST + 1):
                ws.cell(row=r, column=j).fill = FILL_GROUP
            r += 1
            last_group = e["group"]
        st = status.get(e["id"], {})
        stale = bool(st.get("formel")) and st.get("formel") != e["formula"]
        vals = dict(e)
        vals["exValue"] = sig(e["exValue"])
        vals["formula_stamp"] = e["formula"]
        vals["in_status"] = "offen" if stale or not st else st.get("status", "offen")
        vals["in_datum"] = None if stale or not st.get("datum") else dt.date.fromisoformat(st["datum"])
        vals["in_durch"] = None if stale else st.get("durch")
        k = st.get("kommentar")
        ans = st.get("antwort")
        pre = (f"[Formel geändert seit Prüfung vom {st.get('datum', '?')}] " if stale else "") + (f"[Antwort: {ans}] " if ans else "")
        vals["in_kommentar"] = (pre + (k or "")).strip() or None
        for _, _, key in COLS:
            c = ws.cell(row=r, column=C[key], value=vals.get(key))
            c.font, c.border = F_BASE, BORDER
            c.alignment = WRAP if key in ("title", "formula", "exSubst", "ref", "exCase", "in_kommentar") else TOP
            if key in INPUT_KEYS:
                c.fill = FILL_INPUT
                c.protection = Protection(locked=False)
        ws.cell(row=r, column=C["in_datum"]).number_format = "DD.MM.YYYY"
        ws.cell(row=r, column=C["exValue"]).number_format = "General"
        r += 1
    last_row = r - 1

    dv = DataValidation(type="list", formula1='"' + ",".join(STATES) + '"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Freigabe", error="Bitte offen, freigegeben oder abgelehnt wählen.")
    dv.add(f"{status_col}{first}:{status_col}{last_row}")
    ws.add_data_validation(dv)
    dcol = get_column_letter(C["in_datum"])
    dd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2024,1,1)", allow_blank=True, showErrorMessage=True,
                        errorTitle="Datum", error="Bitte ein Datum eingeben, z. B. 12.10.2026.")
    dd.add(f"{dcol}{first}:{dcol}{last_row}")
    ws.add_data_validation(dd)

    full = f"A{first}:{get_column_letter(C['exCase'])}{last_row}"
    ws.conditional_formatting.add(full, FormulaRule(formula=[f'${status_col}{first}="freigegeben"'], fill=FILL_OK))
    ws.conditional_formatting.add(full, FormulaRule(formula=[f'${status_col}{first}="abgelehnt"'], fill=FILL_BAD))

    ws.column_dimensions[get_column_letter(C["formula_stamp"])].hidden = True
    ws.freeze_panes = ws.cell(row=first, column=2)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(LAST)}{last_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{head_row}:{head_row}"
    # Blattschutz ohne Kennwort: nur die gelben Eingabezellen sind änderbar; Filtern, Sortieren und Spaltenbreiten bleiben frei.
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def write_guide(ws, cat: dict) -> None:
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    rows: list[tuple[str, str, str]] = [
        ("title", f"Anleitung Formelfreigabe eng-core {cat['coreVersion']}", ""),
        ("", "", ""),
        ("h", "Ablauf", ""),
        ("kv", "1", "Blatt »Freigabe« öffnen. Jede Zeile ist eine Formel des Rechenkerns, gruppiert wie im FORMELKATALOG.md."),
        ("kv", "2", "Formel prüfen: Formel gegen Quelle, dann das Beispiel nachrechnen (Taschenrechner oder Gegenrechnung). Das Beispiel stammt aus dem Rechenlauf in Spalte »Beispielfall«."),
        ("kv", "3", "In den gelben Spalten eintragen: Freigabe (Auswahlliste), Datum, Prüfer, bei Ablehnung Kommentar oder Verbesserungsvorschlag."),
        ("kv", "4", "Datei speichern (Dateiname egal) und zurückgeben. Teilweise ausgefüllte Dateien sind in Ordnung."),
        ("kv", "5", "Übernahme ins Repo: python3 scripts/formelfreigabe.py import <datei.xlsx>  → formelstatus.json und FORMELKATALOG.md werden aktualisiert, abgelehnte Formeln werden aufgelistet und im Kern korrigiert."),
        ("", "", ""),
        ("h", "Legende", ""),
        ("in", "Gelbe Zellen", "Eingabe. Nur diese Spalten werden beim Import gelesen. Alle anderen Spalten kommen aus dem Rechenkern; Änderungen dort werden ignoriert."),
        ("ok", "Grüne Zeile", "freigegeben: Formel, Quelle und Beispiel geprüft und korrekt."),
        ("bad", "Rote Zeile", "abgelehnt: Fehler oder Zweifel. Bitte im Kommentar beschreiben, was falsch ist oder wie es richtig wäre (mit Quelle, wenn möglich)."),
        ("kv", "offen", "noch nicht geprüft (Standard)."),
        ("kv", "Pflichtfelder", "freigegeben braucht Datum und Prüfer, sonst wird die Zeile beim Import nicht übernommen. abgelehnt braucht einen Kommentar."),
        ("kv", "Formel ändert sich", "Eine Freigabe gilt nur für den geprüften Formeltext. Ändert sich die Formel im Kern, steht sie im Katalog auf »geändert, neu prüfen« und erscheint im nächsten Export wieder als offen."),
        ("", "", ""),
        ("h", "Beispielzeile (so sieht eine ausgefüllte Zeile aus)", ""),
    ]
    r = 1
    for kind, a, b in rows:
        ca, cb = ws.cell(row=r, column=1, value=a or None), ws.cell(row=r, column=2, value=b or None)
        ca.font, cb.font = F_BASE, F_BASE
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ca.alignment = Alignment(vertical="top")
        if kind == "title":
            ca.font = F_TITLE
        elif kind == "h":
            ca.font = F_BOLD
        elif kind == "in":
            ca.fill = FILL_INPUT
        elif kind == "ok":
            ca.fill = FILL_OK
        elif kind == "bad":
            ca.fill = FILL_BAD
        r += 1
    example = [
        ("ID", "B-M-EQ"), ("Formel", "M_max = q·L²/8"), ("Beispiel eingesetzt", "2·1200²/8 = 360000 N·mm"),
        ("Freigabe", "freigegeben"), ("Datum", "12.10.2026"), ("Prüfer", "S. Frotscher"), ("Kommentar / Verbesserung", "(leer, weil freigegeben)"),
    ]
    for a, b in example:
        ws.cell(row=r, column=1, value=a).font = F_BOLD
        c = ws.cell(row=r, column=2, value=b)
        c.font = F_BASE
        if a in ("Freigabe", "Datum", "Prüfer", "Kommentar / Verbesserung"):
            c.fill = FILL_INPUT
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Beispiel Ablehnung").font = F_BOLD
    ws.cell(row=r, column=2, value="Freigabe »abgelehnt«, Kommentar: »Nach Roark Tab. 8.1 Fall 2a gilt f = q·L⁴/(8·E·I); bitte Quelle der Abweichung prüfen.«").font = F_BASE
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True)
    r += 2
    ws.cell(row=r, column=1, value="Hinweise").font = F_BOLD
    r += 1
    for t in (
        "Einheiten intern: N, mm, MPa, kg/m³. Eingaben im Werkzeug: F und N in kN, q in kN/m (= N/mm), T in N·m (im Beispiel in N·mm umgerechnet).",
        "In den Durchbiegungsbeispielen steht E·I als ein Produkt (z. B. 1,941·10¹⁰ N·mm² = 70000 MPa · 2,773·10⁵ mm⁴).",
        "Zahlen im Beispiel sind auf 4 Stellen gerundet angezeigt; das Ergebnis ist auf 6 signifikante Stellen gerundet.",
        "Alle Formeln sind zusätzlich automatisch gegengerechnet (verification/pruefung.py, unabhängige Python-Rechnung). Die Freigabe hier ist die fachliche Prüfung durch einen Menschen.",
    ):
        c = ws.cell(row=r, column=2, value=t)
        c.font, c.alignment = F_BASE, Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Gültigkeit und Annahmen").font = F_BOLD
    r += 1
    for t in (
        "Linear-elastisch, kleine Verformungen, Bernoulli-Balken, statische Last, isotroper Werkstoff.",
        "Torsion Hohlprofile nach Bredt (t ≪ Abmessung); Rechteck voll nach Roark Tab. 10.1 (Näherung, < 0,2 % gegen Saint-Venant-Reihe).",
        "Vergleichsspannung GEH an der Randfaser; τ_Q separat (Maximum an der Nulllinie).",
        "Knicken: Euler für λ ≥ λ_T = π·√(2E/R_p), darunter Johnson-Parabel; keine Imperfektionen, keine Normknickkurven.",
        "Auslegungswert standardmäßig unterer Wert der Datenbankzeile; Normhöchstwerte liefern keinen Auslegungswert.",
        "Bemessen: Bisektion im log-Raum, 60 Schritte, Raster 1 mm (Wandstärke 0,5 mm).",
    ):
        c = ws.cell(row=r, column=2, value=t)
        c.font, c.alignment = F_BASE, Alignment(wrap_text=True, vertical="top")
        r += 1


# --------------------------------------------------------------------------------------------- Import

def import_xlsx(path: Path, dry_run: bool, fill_datum: str | None = None, fill_durch: str | None = None) -> int:
    cat = load_catalog()
    by_id = {e["id"]: e for e in cat["entries"]}
    wb = load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"Blatt »{SHEET}« fehlt in {path.name}")
    ws = wb[SHEET]
    head = next((r for r in range(1, 30) if ws.cell(row=r, column=1).value == "ID"), None)
    if head is None:
        sys.exit("Kopfzeile mit »ID« nicht gefunden")
    names = {str(ws.cell(row=head, column=j).value or "").strip(): j for j in range(1, ws.max_column + 1)}
    col = {key: names.get(name) for name, _, key in COLS}
    for need in ("id", "in_status", "in_datum", "in_durch", "in_kommentar"):
        if not col.get(need):
            sys.exit(f"Spalte »{COLS[C[need] - 1][0]}« fehlt")

    new: dict[str, dict] = {}
    warn: list[str] = []
    seen: set[str] = set()
    for r in range(head + 1, ws.max_row + 1):
        fid = ws.cell(row=r, column=col["id"]).value
        if not fid or fid not in by_id:
            if fid and str(fid).strip() and str(fid) not in {e["groupTitle"] for e in cat["entries"]}:
                warn.append(f"Zeile {r}: unbekannte ID {fid!r}, ignoriert")
            continue
        if fid in seen:
            warn.append(f"Zeile {r}: {fid} doppelt, zweiter Eintrag ignoriert")
            continue
        seen.add(fid)
        e = by_id[fid]
        get = lambda k: ws.cell(row=r, column=col[k]).value if col.get(k) else None
        stat = str(get("in_status") or "offen").strip().lower()
        # automatisch vorangestellte Hinweise ([Formel geändert …], [Antwort: …], [umgesetzt …]) nicht erneut speichern
        kommentar = re.sub(r"^(\[[^\]]*\]\s*)+", "", str(get("in_kommentar") or "").strip()).strip() or None
        durch = str(get("in_durch") or "").strip() or None
        try:
            datum = parse_date(get("in_datum"))
        except ValueError as ex:
            warn.append(f"{fid}: {ex}, Zeile nicht übernommen")
            continue
        if stat not in STATES:
            warn.append(f"{fid}: Freigabe »{stat}« unbekannt, Zeile nicht übernommen")
            continue
        stamp = get("formula_stamp")
        if stamp is not None and stamp != e["formula"]:
            warn.append(f"{fid}: Formel hat sich seit dem Export geändert ({stamp!r} → {e['formula']!r}), Zeile nicht übernommen, bitte neu exportieren")
            continue
        filled = []
        if stat == "freigegeben" and not datum and fill_datum:
            datum = fill_datum; filled.append("Datum")
        if stat == "freigegeben" and not durch and fill_durch:
            durch = fill_durch; filled.append("Prüfer")
        if stat == "freigegeben" and (not datum or not durch):
            warn.append(f"{fid}: freigegeben ohne Datum oder Prüfer, Zeile nicht übernommen")
            continue
        if stat == "abgelehnt" and not kommentar:
            warn.append(f"{fid}: abgelehnt ohne Kommentar")
        if stat == "offen" and not kommentar:
            continue
        prev = cat["status"].get(fid, {})
        erledigt = None
        if stat == "freigegeben" and kommentar and kommentar == prev.get("kommentar"):
            erledigt, kommentar = kommentar, None  # Einwand der Vorrunde, mit dieser Freigabe erledigt
        if filled:
            vermerk = f"{' und '.join(filled)} beim Import ergänzt (Freigabe vom Prüfer außerhalb der Excel bestätigt)"
            kommentar = f"{kommentar} · {vermerk}" if kommentar else vermerk
            warn.append(f"{fid}: {vermerk}")
        rec: dict = {"status": stat}
        if datum:
            rec["datum"] = datum
        if durch:
            rec["durch"] = durch
        if kommentar:
            rec["kommentar"] = kommentar
        if erledigt:
            rec["erledigt"] = erledigt
        if prev.get("antwort"):
            rec["antwort"] = prev["antwort"]  # Antwort des Kern-Verantwortlichen bleibt erhalten
        rec["formel"] = e["formula"]
        new[fid] = rec

    old = cat["status"]
    missing = [i for i in by_id if i not in seen]
    if missing:
        warn.append(f"{len(missing)} Formeln fehlen in der Excel (Teilexport oder älterer Export): {', '.join(missing[:10])}{' …' if len(missing) > 10 else ''}; ihr bisheriger Status bleibt erhalten")
        for i in missing:
            if i in old:
                new[i] = old[i]

    out = {k: new[k] for k in sorted(new)}
    cnt = {s: sum(1 for v in out.values() if v["status"] == s) for s in STATES}
    changed = sorted(k for k in set(out) | set(old) if out.get(k) != old.get(k))
    print(f"{path.name}: {len(seen)} Formeln gelesen · {cnt['freigegeben']} freigegeben · {cnt['abgelehnt']} abgelehnt · {len(changed)} Änderungen gegenüber formelstatus.json")
    for w in warn:
        print(f"  Warnung: {w}")
    rejected = [(k, v) for k, v in out.items() if v["status"] == "abgelehnt" or (v["status"] == "offen" and v.get("kommentar"))]
    if rejected:
        print("\nEinwände und Verbesserungen:")
        for k, v in rejected:
            print(f"  {k} [{v['status']}] {v.get('durch', '')} {v.get('datum', '')}: {v.get('kommentar', '(kein Kommentar)')}")
            print(f"      Formel: {by_id[k]['formula']}")
    if dry_run:
        print("\n--dry-run: nichts geschrieben")
        return 0
    STATUS_FILE.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    res = subprocess.run(["npx", "--no-install", "tsx", "scripts/formelkatalog.ts"], cwd=ROOT, capture_output=True, text=True)
    print(("\n" + res.stdout.strip()) if res.returncode == 0 else f"\nFORMELKATALOG.md konnte nicht erzeugt werden:\n{res.stderr}")
    return res.returncode


def main() -> None:
    ap = argparse.ArgumentParser(description="Formelfreigabe über Excel (Export/Import)")
    sub = ap.add_subparsers(dest="cmd", required=True)
    ex = sub.add_parser("export", help="Freigabe-Excel aus dem Rechenkern erzeugen")
    ex.add_argument("--out", type=Path, default=None)
    ex.add_argument("--offen", action="store_true", help="nur Formeln ohne gültige Freigabe exportieren")
    ex.add_argument("--modul", default=None, help="nur diese Module (Komma) oder ein Werkzeug (werkbank, fachtools)")
    im = sub.add_parser("import", help="Ausgefüllte Excel nach formelstatus.json übernehmen")
    im.add_argument("xlsx", type=Path)
    im.add_argument("--dry-run", action="store_true")
    im.add_argument("--ergaenze-datum", default=None, help="Datum für freigegebene Zeilen ohne Datum (TT.MM.JJJJ)")
    im.add_argument("--ergaenze-pruefer", default=None, help="Prüfer für freigegebene Zeilen ohne Prüfer")
    a = ap.parse_args()
    if a.cmd == "export":
        ver = json.loads((ROOT / "package.json").read_text(encoding="utf-8"))["version"]
        mods = resolve_modules(a.modul)
        tag = f"_{a.modul.replace(',', '-')}" if mods is not None else ""
        name = f"Formelfreigabe_eng-core_{ver}{tag}{'_offen' if a.offen else ''}.xlsx"
        export((a.out or ROOT / "freigabe" / name).resolve(), a.offen, mods)
    else:
        sys.exit(import_xlsx(a.xlsx.resolve(), a.dry_run, parse_date(a.ergaenze_datum), a.ergaenze_pruefer))


if __name__ == "__main__":
    main()
